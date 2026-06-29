# -*- coding: utf-8 -*-
"""
FitBook：指定「球敘」場館之場次；僅寫入「已報名人數 > 0」者。

會員：須登入 Cookie（cookie_header 或 FITBOOK_COOKIE），自課程頁解析「已預約會員」
之暱稱與頭像網址；每位會員一列，頭像可為 Excel 內嵌圖或 Google 試算表 =IMAGE(url)。

異動歷史（v2 設計）：
- 比對基準存在試算表隱藏分頁 `_scan_baseline`（本機與雲端共用同一份基準）
- 每週一 08:00 由 GitHub Actions `FitBook Weekly Reset` 清空歷史、重設基準
- 一般掃描：與基準比對，僅將「新增／刪除」列插入「掃描歷史」最上方
- 任一場解析失敗（Cookie/網路/網頁改版）→ 該次不寫歷史、不更新基準
- 比對僅含「場次日期 >= 今天」場次；無有效基準時當作首次執行
- 本機若無 Google Sheets，退回使用 last_scan_state.json

輸出：若 config 設定 google_sheet_id 且服務帳戶 JSON 存在，寫入 Google 試算表；否則寫本機 Excel。
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import tempfile
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from PIL import Image as PILImage

import pandas as pd
import requests
from bs4 import BeautifulSoup
from requests.exceptions import ChunkedEncodingError, ConnectionError, Timeout

CONFIG_PATH = Path(__file__).resolve().parent / "config.json"
LAST_SCAN_STATE_PATH = Path(__file__).resolve().parent / "last_scan_state.json"
QUICK_SCAN_TICK_PATH = Path(__file__).resolve().parent / "quick_scan_tick.json"
HISTORY_SHEET_NAME = "掃描歷史"
OUTPUT_COLUMNS = [
    "掃描時間",
    "場館標籤",
    "場次日期",
    "星期",
    "時段",
    "已報名人數",
    "開放名額",
    "剩餘名額",
    "會員暱稱",
    "頭像",
    "預約頁面",
]

HISTORY_SHEET_COLUMNS = ["掃描時間", "異動類型"] + [
    c for c in OUTPUT_COLUMNS if c != "掃描時間"
]
# 通知用：含課程名稱；寫入「掃描歷史」試算表時不寫此欄（見 google_sheets_export）
HISTORY_COLUMNS = HISTORY_SHEET_COLUMNS + ["課程名稱"]

# 不含掃描時間：避免比對檔 meta 每次被覆寫，歷史列的掃描時間僅在寫入當下設定一次
SESSION_META_COLUMNS = [
    c
    for c in OUTPUT_COLUMNS
    if c not in ("會員暱稱", "頭像", "掃描時間")
] + ["課程名稱"]

SCAN_STATE_VERSION = 2
# GitHub Actions cache 版本；變更時舊比對檔不會再被還原（避免誤判刪除）
SCAN_STATE_CACHE_KEY = "fitbook-last-scan-v2"


def load_config() -> dict[str, Any]:
    with CONFIG_PATH.open(encoding="utf-8") as f:
        return json.load(f)


def _scan_timezone(cfg: dict[str, Any]) -> ZoneInfo:
    name = (cfg.get("timezone") or "Asia/Taipei").strip()
    try:
        return ZoneInfo(name)
    except Exception:
        return ZoneInfo("Asia/Taipei")


def _normalize_cookie_header(raw: str) -> str:
    """整理 Cookie 字串，避免 Secret 含換行或 'Cookie:' 前綴導致 requests InvalidHeader。"""
    s = (raw or "").strip()
    if not s:
        return ""
    s = s.replace("\r", "").replace("\n", "").strip()
    if s.lower().startswith("cookie:"):
        s = s.split(":", 1)[1].strip()
    return s


def effective_cookie(cfg: dict[str, Any]) -> str:
    c = _normalize_cookie_header(cfg.get("cookie_header") or "")
    if c:
        return c
    env_name = (cfg.get("cookie_env") or "FITBOOK_COOKIE").strip()
    if env_name:
        return _normalize_cookie_header(os.environ.get(env_name) or "")
    return ""


def apply_session_auth(session: requests.Session, cookie: str) -> None:
    """設定 Cookie；若有 XSRF-TOKEN 一併帶入 Laravel 常用標頭。"""
    if not cookie:
        return
    session.headers["Cookie"] = cookie
    for part in cookie.split(";"):
        part = part.strip()
        if part.startswith("XSRF-TOKEN="):
            from urllib.parse import unquote

            session.headers["X-XSRF-TOKEN"] = unquote(part.split("=", 1)[1])
            break


def extract_json_after(html: str, needle: str) -> Any:
    i = html.find(needle)
    if i < 0:
        raise ValueError(f"找不到標記: {needle!r}")
    i += len(needle)
    while i < len(html) and html[i].isspace():
        i += 1
    if i >= len(html) or html[i] != "[":
        raise ValueError("預期為 JSON 陣列")
    depth = 0
    start = i
    for j in range(i, len(html)):
        ch = html[j]
        if ch == "[":
            depth += 1
        elif ch == "]":
            depth -= 1
            if depth == 0:
                return json.loads(html[start : j + 1])
    raise ValueError("JSON 陣列未閉合")


def fetch_home_html(cfg: dict[str, Any], session: requests.Session) -> str:
    url = f"{cfg['base_url']}{cfg['home_path']}"
    r = _session_get_with_retry(session, url)
    r.encoding = r.apparent_encoding or "utf-8"
    return r.text


def parse_course_templates(html: str) -> list[dict[str, Any]]:
    return extract_json_after(html, "let courseTemplates = ")


def course_name_without_venue(template_name: str, venue_label: str) -> str:
    """模板名稱常含「課程(場館)」，拆出純課程名避免通知重複顯示場館。"""
    name = (template_name or "").strip()
    venue = (venue_label or "").strip()
    if not name or not venue or name == venue:
        return name
    for open_p, close_p in (("(", ")"), ("（", "）")):
        suffix = f"{open_p}{venue}{close_p}"
        if name.endswith(suffix):
            return name[: -len(suffix)].strip() or name
    return name


def template_matches_venue(
    name: str,
    venues: list[dict[str, Any]],
    ball_you_keyword: str,
) -> tuple[bool, str | None]:
    if ball_you_keyword and ball_you_keyword not in name:
        return False, None
    for v in venues:
        for m in v["match"]:
            if m in name:
                return True, v["label"]
    return False, None


def fetch_template_courses(
    cfg: dict[str, Any],
    session: requests.Session,
    template_id: int,
) -> dict[str, Any]:
    url = (
        f"{cfg['base_url']}/{cfg['store_path']}/course_template/"
        f"{cfg['place_id']}?id={template_id}"
    )
    r = _session_get_with_retry(session, url)
    return r.json()


def _is_login_wall(html: str) -> bool:
    if "註冊登入" in html and "使用LINE帳號登入" in html:
        return True
    t = re.search(r"<title>([^<]+)</title>", html, re.I)
    if t and "註冊登入" in t.group(1):
        return True
    return False


def _looks_like_person_name(s: str) -> bool:
    s = s.strip()
    if not s or len(s) < 2 or len(s) > 24:
        return False
    if re.fullmatch(r"[\u4e00-\u9fff\u00b7\u2022\u30fb·\sA-Za-z0-9]+", s):
        return True
    return False


def _normalize_img_url(src: str | None) -> str | None:
    if not src or not str(src).strip():
        return None
    u = str(src).strip()
    if u.startswith("//"):
        u = "https:" + u
    if u.startswith("http"):
        return u
    return None


def _parse_fitbook_yiyu_huiyuan(soup: BeautifulSoup) -> list[tuple[str, str | None]]:
    """FitBook〈已預約會員〉：(暱稱, 頭像 URL)，順序與網頁一致，允許暱稱重複。"""
    out: list[tuple[str, str | None]] = []
    for h2 in soup.find_all("h2"):
        if "已預約會員" not in h2.get_text():
            continue
        container = h2.find_next_sibling("div")
        if not container:
            continue
        for block in container.select("div.d-inline-block.w-4em"):
            img = block.find("img", alt=True)
            ptag = block.find("p", class_=lambda c: bool(c) and "truncate" in c)
            alt = (img.get("alt") or "").strip() if img else ""
            ptxt = ptag.get_text(strip=True) if ptag else ""
            url = _normalize_img_url(img.get("src") if img else None)
            if ptxt:
                chosen = ptxt
            else:
                chosen = alt
            if chosen:
                out.append((chosen, url))
    return out


def parse_members_with_avatars_from_html(
    html: str, teacher_name: str | None, extra_selectors: list[str] | None
) -> list[tuple[str, str | None]]:
    """從課程頁擷取會員暱稱與頭像網址（無圖則為 None）。"""
    if _is_login_wall(html):
        return []

    soup = BeautifulSoup(html, "html.parser")

    fb = _parse_fitbook_yiyu_huiyuan(soup)
    if fb:
        if teacher_name:
            tn = teacher_name.strip()
            fb = [(n, u) for n, u in fb if n != tn]
        return fb

    names: list[str] = []
    for sel in extra_selectors or []:
        try:
            for el in soup.select(sel):
                t = el.get_text(strip=True)
                if _looks_like_person_name(t) and t not in names:
                    names.append(t)
        except Exception:
            continue

    for tr in soup.select("table tbody tr"):
        tds = tr.find_all("td")
        if not tds:
            continue
        for td in tds[:2]:
            t = td.get_text(strip=True)
            if _looks_like_person_name(t):
                if teacher_name and t == teacher_name.strip():
                    continue
                if t not in names:
                    names.append(t)
            break

    for script in soup.find_all("script"):
        txt = script.string
        if not txt or len(txt) < 20:
            continue
        for m in re.finditer(
            r'"(?:member_name|memberName|nickname|reserve_name|member_nickname)"\s*:\s*"([^"\\]+)"',
            txt,
        ):
            val = m.group(1).strip()
            if _looks_like_person_name(val) and val not in names:
                names.append(val)

    for tag in soup.find_all(True, attrs={"data-member-name": True}):
        val = tag.get("data-member-name", "").strip()
        if _looks_like_person_name(val) and val not in names:
            names.append(val)

    if teacher_name:
        tn = teacher_name.strip()
        names = [n for n in names if n != tn]

    seen: set[str] = set()
    out_n: list[str] = []
    for n in names:
        if n not in seen:
            seen.add(n)
            out_n.append(n)
    return [(n, None) for n in out_n]


def fetch_member_course_html(
    cfg: dict[str, Any],
    session: requests.Session,
    course_id: int,
    place_id: int,
) -> str:
    url = f"{cfg['base_url']}/{cfg['store_path']}/member/course/{course_id}/{place_id}"
    headers = {
        "Referer": f"{cfg['base_url']}{cfg.get('home_path') or ''}",
        "Accept": "text/html,application/xhtml+xml;q=0.9,*/*;q=0.8",
    }
    r = _session_get_with_retry(session, url, headers=headers)
    r.encoding = r.apparent_encoding or "utf-8"
    return r.text


def session_members_snapshot(df: pd.DataFrame) -> dict[str, list[str]]:
    """以預約頁面 URL 為場次鍵，值為該場會員暱稱列表（順序與列相同，含重複）。"""
    if df.empty or "預約頁面" not in df.columns or "會員暱稱" not in df.columns:
        return {}
    out: dict[str, list[str]] = {}
    for url, grp in df.groupby("預約頁面", sort=False):
        out[str(url)] = list(grp["會員暱稱"].astype(str))
    return out


def _json_safe_value(val: Any) -> Any:
    """將 pandas/numpy 純量轉成 JSON 可序列化的 Python 型別。"""
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    if hasattr(val, "item"):
        try:
            return val.item()
        except (ValueError, AttributeError):
            pass
    if isinstance(val, (datetime, date)):
        return val.isoformat()
    return val


def _members_map_from_state(state: dict[str, Any]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for k, v in state.items():
        if isinstance(v, dict) and "members" in v:
            out[str(k)] = list(v["members"])
        elif isinstance(v, list):
            out[str(k)] = list(v)
    return out


def _session_date_from_meta(meta: dict[str, Any]) -> date | None:
    return _parse_session_date_val(meta.get("場次日期"))


def filter_state_active_sessions(
    state: dict[str, dict[str, Any]], today: date
) -> dict[str, dict[str, Any]]:
    """只保留場次日期 >= today 的場次，避免已過期活動被誤判為「刪除」。"""
    out: dict[str, dict[str, Any]] = {}
    for url, entry in state.items():
        if not isinstance(entry, dict):
            continue
        meta = entry.get("meta") or {}
        sdate = _session_date_from_meta(meta)
        if sdate is not None and sdate < today:
            continue
        out[str(url)] = entry
    return out


def _session_meta_from_group(grp: pd.DataFrame) -> dict[str, Any]:
    if grp.empty:
        return {}
    first = grp.iloc[0]
    meta: dict[str, Any] = {}
    for c in SESSION_META_COLUMNS:
        if c in grp.columns:
            meta[c] = _json_safe_value(first[c])
    return meta


def _avatar_url_map_from_df(
    df: pd.DataFrame, avatar_urls: list[str | None] | None
) -> dict[tuple[str, str], str]:
    """(預約頁面, 會員暱稱) → 頭像 URL。"""
    m: dict[tuple[str, str], str] = {}
    if df.empty or not avatar_urls or len(avatar_urls) != len(df):
        return m
    for pos, (_, row) in enumerate(df.iterrows()):
        page = str(row.get("預約頁面", "")).strip()
        name = str(row.get("會員暱稱", "")).strip()
        if not page or not name or name.startswith("（"):
            continue
        u = avatar_urls[pos]
        if u:
            m[(page, name)] = str(u).strip()
    return m


def _is_placeholder_member_name(name: str) -> bool:
    """括號開頭為解析失敗說明，不可當真實會員比對。"""
    return str(name or "").strip().startswith("（")


def _members_list_from_group(grp: pd.DataFrame) -> list[str]:
    out: list[str] = []
    for n in grp["會員暱稱"].astype(str):
        n = n.strip()
        if n and not _is_placeholder_member_name(n):
            out.append(n)
    return out


def _count_placeholder_sessions(df: pd.DataFrame) -> int:
    """有報名但會員欄僅為錯誤說明文字的場次數。"""
    if df.empty or "會員暱稱" not in df.columns or "預約頁面" not in df.columns:
        return 0
    n = 0
    for _, grp in df.groupby("預約頁面", sort=False):
        names = _members_list_from_group(grp)
        if not names and any(
            _is_placeholder_member_name(x) for x in grp["會員暱稱"].astype(str)
        ):
            n += 1
    return n


def _member_avatars_for_group(
    grp: pd.DataFrame, avatar_map: dict[tuple[str, str], str], url: str
) -> dict[str, str]:
    out: dict[str, str] = {}
    for name in grp["會員暱稱"].astype(str):
        name = name.strip()
        if not name or name.startswith("（"):
            continue
        u = avatar_map.get((str(url), name), "")
        if u:
            out[name] = u
    return out


def session_state_from_dataframe(
    df: pd.DataFrame, avatar_urls: list[str | None] | None = None
) -> dict[str, dict[str, Any]]:
    """場次鍵 → {members, meta, member_avatars}，供比對與刪除列還原場次欄位。"""
    if df.empty or "預約頁面" not in df.columns or "會員暱稱" not in df.columns:
        return {}
    avatar_map = _avatar_url_map_from_df(df, avatar_urls)
    out: dict[str, dict[str, Any]] = {}
    for url, grp in df.groupby("預約頁面", sort=False):
        first = grp.iloc[0]
        meta: dict[str, Any] = {}
        for c in SESSION_META_COLUMNS:
            if c in grp.columns:
                meta[c] = _json_safe_value(first[c])
        out[str(url)] = {
            "members": _members_list_from_group(grp),
            "meta": meta,
            "member_avatars": _member_avatars_for_group(grp, avatar_map, str(url)),
        }
    return out


def load_last_scan_state() -> dict[str, dict[str, Any]]:
    if not LAST_SCAN_STATE_PATH.is_file():
        return {}
    try:
        raw = json.loads(LAST_SCAN_STATE_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError, TypeError):
        return {}
    if not isinstance(raw, dict):
        return {}
    if raw.get("version") == SCAN_STATE_VERSION and isinstance(raw.get("sessions"), dict):
        sessions = raw["sessions"]
        return {
            str(k): v
            for k, v in sessions.items()
            if isinstance(v, dict) and "members" in v
        }
    # v1：僅 { url: [暱稱, ...] }
    return {
        str(k): {"members": list(v), "meta": {}}
        for k, v in raw.items()
        if isinstance(v, list)
    }


def save_last_scan_state(state: dict[str, dict[str, Any]]) -> None:
    LAST_SCAN_STATE_PATH.write_text(
        json.dumps(
            {"version": SCAN_STATE_VERSION, "sessions": state},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def snapshot_has_member_change(
    previous: dict[str, dict[str, Any]], current: dict[str, dict[str, Any]]
) -> bool:
    """任一场次暱稱人次（Counter）不同、場次新增或消失，皆視為有異動。"""
    prev_m = _members_map_from_state(previous)
    curr_m = _members_map_from_state(current)
    keys = set(prev_m) | set(curr_m)
    for k in keys:
        if Counter(prev_m.get(k, [])) != Counter(curr_m.get(k, [])):
            return True
    return False


def _multiset_added_removed(
    previous: list[str], current: list[str]
) -> tuple[list[str], list[str]]:
    pc, cc = Counter(previous), Counter(current)
    added: list[str] = []
    removed: list[str] = []
    for name in set(pc) | set(cc):
        delta = cc[name] - pc[name]
        if delta > 0:
            added.extend([name] * delta)
        delta = pc[name] - cc[name]
        if delta > 0:
            removed.extend([name] * delta)
    return added, removed


def _take_member_row(
    grp: pd.DataFrame, name: str, used: set[Any]
) -> pd.Series | None:
    for idx, row in grp.iterrows():
        if idx in used:
            continue
        if str(row.get("會員暱稱", "")) == name:
            used.add(idx)
            return row
    return None


def build_history_diff_chunk(
    df: pd.DataFrame,
    previous: dict[str, dict[str, Any]],
    current: dict[str, dict[str, Any]],
    columns: list[str],
    today: date | None = None,
    avatar_urls: list[str | None] | None = None,
) -> pd.DataFrame:
    """僅產生「新增」「刪除」列；不含場次人數等純欄位更新。已過期場次不參與比對。"""
    if today is None:
        today = date.today()
    previous = filter_state_active_sessions(previous, today)
    current = filter_state_active_sessions(current, today)

    scan_at = ""
    if not df.empty and "掃描時間" in df.columns:
        scan_at = str(df["掃描時間"].iloc[0])

    records: list[dict[str, Any]] = []
    urls = set(previous) | set(current)
    avatar_map = _avatar_url_map_from_df(df, avatar_urls)

    for url in urls:
        prev_entry = previous.get(url, {})
        curr_entry = current.get(url, {})
        prev_m = (
            list(prev_entry.get("members", []))
            if isinstance(prev_entry, dict)
            else []
        )
        curr_m = (
            list(curr_entry.get("members", []))
            if isinstance(curr_entry, dict)
            else []
        )
        added, removed = _multiset_added_removed(prev_m, curr_m)
        if not added and not removed:
            continue

        prev_meta = (
            dict(prev_entry.get("meta") or {})
            if isinstance(prev_entry, dict)
            else {}
        )
        curr_meta = (
            dict(curr_entry.get("meta") or {})
            if isinstance(curr_entry, dict)
            else {}
        )
        curr_grp = (
            df[df["預約頁面"].astype(str) == str(url)]
            if not df.empty and "預約頁面" in df.columns
            else pd.DataFrame()
        )
        used: set[Any] = set()

        for name in added:
            base: dict[str, Any]
            row = _take_member_row(curr_grp, name, used) if not curr_grp.empty else None
            if row is not None:
                base = {
                    c: row[c]
                    for c in row.index
                    if c in columns and c != "掃描時間"
                }
            else:
                base = {**curr_meta, "預約頁面": url, "會員暱稱": name}
            base["掃描時間"] = scan_at
            base["異動類型"] = "新增"
            base["頭像"] = avatar_map.get((str(url), name), "")
            records.append(base)

        prev_avatars = (
            prev_entry.get("member_avatars")
            if isinstance(prev_entry, dict)
            else {}
        )
        if not isinstance(prev_avatars, dict):
            prev_avatars = {}

        for name in removed:
            if not curr_grp.empty:
                base = _session_meta_from_group(curr_grp)
            else:
                base = dict(prev_meta)
            base["預約頁面"] = url
            base["會員暱稱"] = name
            base["頭像"] = str(prev_avatars.get(name, "") or "")
            base.pop("掃描時間", None)
            base["掃描時間"] = scan_at
            base["異動類型"] = "刪除"
            records.append(base)

    if not records:
        return pd.DataFrame(columns=columns)
    return pd.DataFrame(records).reindex(columns=columns).fillna("")


def load_history_dataframe(
    out_path: Path, columns: list[str], cfg: dict[str, Any]
) -> pd.DataFrame:
    from google_sheets_export import google_sheets_enabled, load_history_from_gsheet

    if google_sheets_enabled(cfg):
        return load_history_from_gsheet(cfg, columns)
    if not out_path.is_file():
        return pd.DataFrame(columns=columns)
    try:
        hist = pd.read_excel(out_path, sheet_name=HISTORY_SHEET_NAME)
    except (ValueError, OSError):
        return pd.DataFrame(columns=columns)
    if "異動類型" not in hist.columns:
        return pd.DataFrame(columns=columns)
    hist = hist.reindex(columns=columns)
    if "掃描時間" in hist.columns:
        hist["掃描時間"] = hist["掃描時間"].astype(str).str.strip()
    return hist


def _parse_session_date_val(val: Any) -> date | None:
    """從 API 的 date_val 轉成日期；無法解析則回傳 None。"""
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    m = re.match(
        r"^(\d{4})[./\-年](\d{1,2})[./\-月](\d{1,2})(?:日)?",
        s,
    )
    if m:
        y, mo, d_ = int(m[1]), int(m[2]), int(m[3])
        return date(y, mo, d_)
    ts = pd.to_datetime(s, errors="coerce", utc=False)
    if pd.isna(ts):
        return None
    pdt = ts.to_pydatetime()
    return pdt.date()


def _coerce_nonneg_int(x: Any) -> int:
    if x is None or (isinstance(x, (float,)) and pd.isna(x)):
        return 0
    if isinstance(x, (int, float)) and not isinstance(x, bool):
        try:
            v = int(x)
        except (ValueError, OSError, OverflowError):
            return 0
        return v if v >= 0 else 0
    s = str(x).strip()
    if not s or s.lower() in ("nan", "none", "nat"):
        return 0
    try:
        v = int(float(s))
    except (TypeError, ValueError, OverflowError):
        return 0
    return v if v >= 0 else 0


def _as_text_cell(x: Any) -> str:
    if x is None or (isinstance(x, (float,)) and pd.isna(x)):
        return ""
    return str(x).strip()


def _slot_start_minutes(slot: Any) -> int:
    """從「19:30~21:30」等字串取開始時間（分鐘），供時段排序；無法解析排最後。"""
    s = str(slot or "").replace("～", "~").replace("–", "-")
    m = re.search(r"(\d{1,2}):(\d{2})", s)
    if not m:
        return 99_999
    try:
        return int(m.group(1)) * 60 + int(m.group(2))
    except (TypeError, ValueError):
        return 99_999


def _sort_dataframe_by_session_date(df: pd.DataFrame) -> pd.DataFrame:
    """以場次日期遞增排序（日期越接近今天越靠上）；無法解析者排最後。
    同日再依時段開始時間遞增（早的在上）；其後：預約頁面、場館、報名順序。
    """
    if df.empty or "場次日期" not in df.columns:
        return df
    out = df.copy()
    out["_dt"] = pd.to_datetime(out["場次日期"], errors="coerce")
    if "時段" in out.columns:
        out["_slot"] = out["時段"].map(_slot_start_minutes)
    sub: tuple[str, ...] = ("預約頁面", "場館標籤")
    if "__member_order__" in out.columns:
        sub = sub + ("__member_order__",)
    elif "會員暱稱" in out.columns:
        sub = sub + ("會員暱稱",)
    by = ["_dt"]
    if "_slot" in out.columns:
        by.append("_slot")
    by.extend(c for c in sub if c in out.columns)
    out = out.sort_values(by=by, ascending=True, na_position="last").reset_index(
        drop=True
    )
    drop_cols = [c for c in ("_dt", "_slot") if c in out.columns]
    return out.drop(columns=drop_cols)


def reset_scan_baseline(*, clear_history: bool = True) -> None:
    """刪除比對檔，下次執行視為初始版（不追加掃描歷史）。clear_history 僅供呼叫端提示。"""
    if LAST_SCAN_STATE_PATH.is_file():
        LAST_SCAN_STATE_PATH.unlink()
    if QUICK_SCAN_TICK_PATH.is_file():
        try:
            QUICK_SCAN_TICK_PATH.unlink()
        except OSError:
            pass
    _ = clear_history


def quick_scan_enabled(cfg: dict[str, Any]) -> bool:
    if cfg.get("quick_scan_enabled") is True:
        return True
    if cfg.get("quick_scan_enabled") is False:
        return False
    env = (os.environ.get("QUICK_SCAN_ENABLED") or "").strip().lower()
    if env in ("1", "true", "yes"):
        return True
    if env in ("0", "false", "no"):
        return False
    return bool(cfg.get("quick_scan_enabled"))


def quick_scan_interval_minutes(cfg: dict[str, Any]) -> int:
    try:
        n = int(cfg.get("quick_scan_interval_minutes") or 2)
    except (TypeError, ValueError):
        n = 2
    return max(1, min(n, 60))


def _full_scan_every_quick_checks(cfg: dict[str, Any]) -> int:
    try:
        n = int(cfg.get("full_scan_every_quick_checks") or 15)
    except (TypeError, ValueError):
        n = 15
    return max(0, n)


def _load_quick_scan_tick() -> dict[str, Any]:
    if not QUICK_SCAN_TICK_PATH.is_file():
        return {"count": 0}
    try:
        raw = json.loads(QUICK_SCAN_TICK_PATH.read_text(encoding="utf-8"))
        if isinstance(raw, dict):
            return raw
    except (json.JSONDecodeError, OSError, TypeError):
        pass
    return {"count": 0}


def _save_quick_scan_tick(data: dict[str, Any]) -> None:
    QUICK_SCAN_TICK_PATH.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _should_force_periodic_full_scan(cfg: dict[str, Any]) -> bool:
    """每 N 次快掃強制深度掃描一次，避免僅換人但人數相同時漏判。"""
    every = _full_scan_every_quick_checks(cfg)
    if every <= 0:
        return False
    tick = _load_quick_scan_tick()
    count = int(tick.get("count") or 0) + 1
    if count >= every:
        _save_quick_scan_tick({"count": 0})
        return True
    _save_quick_scan_tick({"count": count})
    return False


def _make_http_session(cfg: dict[str, Any]) -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            ),
            "Accept-Language": "zh-TW,zh;q=0.9",
        }
    )
    apply_session_auth(session, effective_cookie(cfg))
    return session


def _session_get_with_retry(
    session: requests.Session,
    url: str,
    *,
    headers: dict[str, str] | None = None,
    timeout: int | float = 60,
    tries: int = 3,
    pause: float = 2.0,
) -> requests.Response:
    """GET 含短暫重試，應付 FitBook 偶發斷線（RemoteDisconnected）。"""
    last_err: Exception | None = None
    for attempt in range(tries):
        try:
            r = session.get(url, headers=headers, timeout=timeout)
            if r.status_code in (429, 502, 503, 504) and attempt < tries - 1:
                time.sleep(pause * (attempt + 1))
                continue
            r.raise_for_status()
            return r
        except (ConnectionError, Timeout, ChunkedEncodingError) as e:
            last_err = e
            if attempt < tries - 1:
                time.sleep(pause * (attempt + 1))
    if last_err is not None:
        raise last_err
    raise RuntimeError(f"GET 失敗：{url}")


def _iter_template_courses(
    cfg: dict[str, Any],
    session: requests.Session,
    *,
    today: date,
) -> list[tuple[str, dict[str, Any]]]:
    """回傳 (模板名稱, 課程 dict) 列表，篩選規則與 run_once 相同。"""
    venues = cfg["venues"]
    ball_kw = cfg.get("only_ball_you_name_contains") or ""
    match_all_templates = bool(cfg.get("match_all_templates"))
    include_zero_reservation = bool(cfg.get("include_zero_reservation_sessions"))

    html = fetch_home_html(cfg, session)
    templates = parse_course_templates(html)
    out: list[tuple[str, dict[str, Any]]] = []
    for t in templates:
        name = t.get("name") or ""
        ok, venue_label = template_matches_venue(name, venues, ball_kw)
        if not ok and not match_all_templates:
            continue
        tid = int(t["id"])
        data = fetch_template_courses(cfg, session, tid)
        time.sleep(0.15)
        for c in data.get("courses") or []:
            try:
                rc_int = int(c.get("reservation_count") or 0)
            except (TypeError, ValueError):
                rc_int = 0
            if rc_int <= 0 and not include_zero_reservation:
                continue
            sdate = _parse_session_date_val(c.get("date_val"))
            if sdate is not None and sdate < today:
                continue
            out.append((str(name), c))
    return out


def fetch_light_session_counts(
    cfg: dict[str, Any],
    session: requests.Session | None = None,
    *,
    today: date | None = None,
) -> dict[str, int]:
    """僅透過首頁 + course_template API 取得各場次已報名人數（不抓會員頁）。"""
    own_session = session is None
    if own_session:
        session = _make_http_session(cfg)
    if today is None:
        today = datetime.now(_scan_timezone(cfg)).date()
    counts: dict[str, int] = {}
    for _, c in _iter_template_courses(cfg, session, today=today):
        url = (c.get("url") or "").strip()
        if not url:
            continue
        try:
            rc_int = int(c.get("reservation_count") or 0)
        except (TypeError, ValueError):
            rc_int = 0
        counts[url] = rc_int
    return counts


def light_counts_changed_since_baseline(
    baseline: dict[str, dict[str, Any]],
    light_counts: dict[str, int],
    today: date,
) -> bool:
    """比對 API 人數與基準會員數；任一場次不同即視為可能有異動。"""
    baseline_active = filter_state_active_sessions(baseline, today)
    urls = set(baseline_active) | set(light_counts)
    for url in urls:
        api_n = int(light_counts.get(url, 0))
        entry = baseline_active.get(url, {})
        if isinstance(entry, dict):
            base_n = len(entry.get("members") or [])
        else:
            base_n = 0
        if api_n != base_n:
            return True
    return False


def run_quick_check(cfg: dict[str, Any]) -> tuple[bool, str]:
    """執行快掃。回傳 (是否需要深度掃描, 說明)。"""
    tz = _scan_timezone(cfg)
    today = datetime.now(tz).date()
    from google_sheets_export import google_sheets_enabled, open_spreadsheet

    sheets_on = google_sheets_enabled(cfg)
    sh = open_spreadsheet(cfg) if sheets_on else None
    baseline = _load_baseline(cfg, sh)
    baseline_compare = filter_state_active_sessions(baseline, today)

    if _should_force_periodic_full_scan(cfg):
        return True, "定期完整掃描（防漏同額換人）"

    if not baseline_compare:
        return True, "尚無比對基準，需完整掃描"

    try:
        counts = fetch_light_session_counts(cfg, today=today)
    except (ConnectionError, Timeout, ChunkedEncodingError, requests.HTTPError) as e:
        return False, f"快掃連線失敗，略過本次（{type(e).__name__}）"
    if light_counts_changed_since_baseline(baseline_compare, counts, today):
        changed = []
        for url in set(baseline_compare) | set(counts):
            api_n = int(counts.get(url, 0))
            entry = baseline_compare.get(url, {})
            base_n = len((entry.get("members") or []) if isinstance(entry, dict) else [])
            if api_n != base_n:
                changed.append(f"{base_n}→{api_n}")
        hint = f"（{len(changed)} 場人數變動）" if changed else ""
        return True, f"偵測到報名人數異動{hint}"

    return False, "無報名人數異動"


def run_once() -> tuple[pd.DataFrame, list[str | None], requests.Session]:
    cfg = load_config()
    venues = cfg["venues"]
    ball_kw = cfg.get("only_ball_you_name_contains") or ""
    match_all_templates = bool(cfg.get("match_all_templates"))
    cookie = effective_cookie(cfg)
    extra_selectors = cfg.get("member_name_css_selectors") or []

    session = _make_http_session(cfg)

    tz = _scan_timezone(cfg)
    now_local = datetime.now(tz)
    scan_at = now_local.strftime("%Y-%m-%d %H:%M:%S")
    today = now_local.date()

    rows: list[dict[str, Any]] = []
    avatar_urls: list[str | None] = []

    for name, c in _iter_template_courses(cfg, session, today=today):
        ok, venue_label = template_matches_venue(
            name, venues, ball_kw
        )
        if not ok and not match_all_templates:
            continue
        if not venue_label:
            venue_label = name.strip() or "未分類"

        rc = c.get("reservation_count")
        try:
            rc_int = int(rc) if rc is not None else 0
        except (TypeError, ValueError):
            rc_int = 0

        teacher_name = c.get("teacher_name") or ""
        members: list[tuple[str, str | None]] = []
        name_note = ""

        if rc_int > 0 and cookie:
            try:
                page_html = fetch_member_course_html(
                    cfg, session, int(c["id"]), int(cfg["place_id"])
                )
                if _is_login_wall(page_html):
                    name_note = (
                        "（Cookie 可能已過期或未登入，課程頁被導向登入頁；"
                        "請重新登入 FitBook 後更新 Cookie）"
                    )
                else:
                    members = parse_members_with_avatars_from_html(
                        page_html, teacher_name, extra_selectors
                    )
                    if not members:
                        name_note = (
                            "（課程頁已登入但未解析到姓名，"
                            "可能網頁改版；可調整 member_name_css_selectors）"
                        )
            except requests.RequestException:
                name_note = "（課程頁請求失敗）"
        elif rc_int > 0:
            name_note = "（請在 config.json 設定 cookie_header 或環境變數 FITBOOK_COOKIE）"

        base_row = {
            "掃描時間": scan_at,
            "課程名稱": course_name_without_venue(name, venue_label),
            "場館標籤": venue_label or "",
            "場次日期": _as_text_cell(c.get("date_val")) or _as_text_cell(
                c.get("date")
            ),
            "星期": _as_text_cell(c.get("day_of_week_val"))
            or _as_text_cell(c.get("day_of_week")),
            "時段": _as_text_cell(c.get("show_time")) or _as_text_cell(
                c.get("time")
            ),
            "已報名人數": int(rc_int),
            "開放名額": _coerce_nonneg_int(c.get("order_count")),
            "剩餘名額": _coerce_nonneg_int(c.get("remain_count")),
            "預約頁面": (c.get("url") or "").strip(),
        }

        if members:
            for mi, (mname, murl) in enumerate(members):
                r = dict(base_row)
                r["會員暱稱"] = mname
                r["頭像"] = ""
                r["__member_order__"] = mi
                rows.append(r)
                avatar_urls.append(murl)
        else:
            r = dict(base_row)
            r["會員暱稱"] = name_note
            r["頭像"] = ""
            r["__member_order__"] = 0
            rows.append(r)
            avatar_urls.append(None)

    df = pd.DataFrame(rows)
    if not df.empty:
        member_order = (
            df.pop("__member_order__") if "__member_order__" in df.columns else None
        )
        keep_cols = [c for c in OUTPUT_COLUMNS if c in df.columns]
        if "課程名稱" in df.columns:
            keep_cols.append("課程名稱")
        df = df[keep_cols]
        if member_order is not None:
            df["__member_order__"] = member_order.values
        n = len(df)
        if len(avatar_urls) == n and n:
            tmp_col = "__avatar_order__"
            df[tmp_col] = avatar_urls
            df = _sort_dataframe_by_session_date(df)
            ser = df.pop(tmp_col)
            avatar_urls = ser.tolist()
        else:
            df = _sort_dataframe_by_session_date(df)
        if "__member_order__" in df.columns:
            df = df.drop(columns=["__member_order__"])
    return df, avatar_urls, session


def _excel_display_width(s: str) -> float:
    """估算儲存格文字在欄寬上的可見長度（中英混排：全形字元約 2、ASCII 約 1）。"""
    w = 0.0
    for ch in str(s):
        if "\u4e00" <= ch <= "\u9fff" or ord(ch) > 0x2E7F:
            w += 2.1
        else:
            w += 1.05
    return max(w, 1.0)


SESSION_MERGE_COLUMNS = [
    "掃描時間",
    "場館標籤",
    "場次日期",
    "星期",
    "時段",
    "已報名人數",
    "開放名額",
    "剩餘名額",
    "預約頁面",
]


def _emptyish_for_merge(x: Any) -> bool:
    if x is None or (isinstance(x, (float,)) and pd.isna(x)):
        return True
    if x == "" or (isinstance(x, str) and not str(x).strip()):
        return True
    s = str(x).strip()
    if s in ("", "None", "nan", "NaT"):
        return True
    return False


def _cell_equal_for_merge(a: Any, b: Any) -> bool:
    if _emptyish_for_merge(a) and _emptyish_for_merge(b):
        return True
    if _emptyish_for_merge(a) and not _emptyish_for_merge(b):
        try:
            if float(b) == 0.0:
                return True
        except (TypeError, ValueError, OverflowError):
            pass
        return False
    if _emptyish_for_merge(b) and not _emptyish_for_merge(a):
        try:
            if float(a) == 0.0:
                return True
        except (TypeError, ValueError, OverflowError):
            pass
        return False
    if a == b:
        return True
    try:
        if float(a) == float(b):
            return True
    except (TypeError, ValueError, OverflowError):
        pass
    return str(a).strip() == str(b).strip()


def _merge_text_norm(x: Any) -> str:
    s = ("" if _emptyish_for_merge(x) else str(x)).replace("～", "~").replace("–", "-")
    s = re.sub(r"\s+", " ", s.strip())
    return s


def _merge_booking_key(x: Any) -> str:
    s = ("" if _emptyish_for_merge(x) else str(x)).strip()
    s = s.split("#", 1)[0].split("?", 1)[0]
    m = re.search(r"/course/(\d+)/(\d+)", s)
    if m:
        return f"{m.group(1)}/{m.group(2)}"
    return s


def _merge_date_equal(a: Any, b: Any) -> bool:
    if _emptyish_for_merge(a) and _emptyish_for_merge(b):
        return True
    d1 = pd.to_datetime(a, errors="coerce")
    d2 = pd.to_datetime(b, errors="coerce")
    if not pd.isna(d1) and not pd.isna(d2) and d1.normalize() == d2.normalize():
        return True
    return _merge_text_norm(a) == _merge_text_norm(b)


def _merge_field_equal(col: str, a: Any, b: Any) -> bool:
    if col == "預約頁面":
        ka, kb = _merge_booking_key(a), _merge_booking_key(b)
        if ka and kb and ka == kb:
            return True
        return _emptyish_for_merge(a) and _emptyish_for_merge(b)
    if col == "場次日期":
        return _merge_date_equal(a, b)
    if col in ("掃描時間", "場館標籤", "星期", "時段"):
        if _cell_equal_for_merge(a, b):
            return True
        return _merge_text_norm(a) == _merge_text_norm(b)
    return _cell_equal_for_merge(a, b)


def _rows_same_session_block(df: pd.DataFrame, i: int, j: int) -> bool:
    for c in SESSION_MERGE_COLUMNS:
        if c not in df.columns:
            return False
        if not _merge_field_equal(c, df.iloc[i][c], df.iloc[j][c]):
            return False
    return True


def _merge_session_info_cells(ws: Any, df: pd.DataFrame) -> None:
    """連續多列若 A～H 與預約頁面均相同，則垂直合併該等欄（不併 會員／頭像 欄）。"""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    if df.empty or len(df) < 2:
        return
    if not all(c in df.columns for c in SESSION_MERGE_COLUMNS):
        return

    col_letters = {
        c: get_column_letter(list(df.columns).index(c) + 1) for c in SESSION_MERGE_COLUMNS
    }

    i = 0
    n = len(df)
    while i < n:
        j = i + 1
        while j < n and _rows_same_session_block(df, i, j):
            j += 1
        if j > i + 1:
            start_excel = i + 2
            end_excel = j + 1
            for c in SESSION_MERGE_COLUMNS:
                letter = col_letters[c]
                mrange = f"{letter}{start_excel}:{letter}{end_excel}"
                try:
                    ws.merge_cells(mrange)
                except (ValueError, OSError, TypeError, KeyError):
                    continue
                for r in range(start_excel + 1, end_excel + 1):
                    cell = ws[f"{letter}{r}"]
                    if not isinstance(cell, MergedCell):
                        cell.value = None
                top = ws[f"{letter}{start_excel}"]
                if not isinstance(top, MergedCell):
                    top.alignment = Alignment(
                        horizontal="left", vertical="center", wrap_text=True
                    )
        i = j


def _autofit_worksheet_columns(ws: Any, df: pd.DataFrame) -> None:
    from openpyxl.utils import get_column_letter

    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        label = str(col)
        max_w = _excel_display_width(label)

        if len(df) == 0:
            ws.column_dimensions[letter].width = min(max(max_w * 1.08 + 2.0, 14.0), 70.0)
            continue

        col_key = label
        for v in df[col].astype(str):
            if col_key == "預約頁面":
                sample = v if len(v) <= 90 else v[:90]
                max_w = max(max_w, _excel_display_width(sample))
            else:
                max_w = max(max_w, _excel_display_width(v))

        if col_key == "預約頁面":
            width = min(max(max_w * 1.05 + 2.5, 18.0), 58.0)
        elif col_key == "會員暱稱":
            width = min(max(max_w * 1.05 + 2.5, 18.0), 72.0)
        elif col_key == "頭像":
            width = 12.0
        else:
            width = min(max(max_w * 1.08 + 2.5, 14.0), 48.0)

        ws.column_dimensions[letter].width = width


def _download_avatar_image(session: requests.Session, url: str) -> bytes | None:
    try:
        r = session.get(
            url,
            timeout=30,
            headers={"Referer": "https://www.fit-book.com.tw/"},
        )
        r.raise_for_status()
        if len(r.content) > 2_500_000:
            return None
        return r.content
    except requests.RequestException:
        return None


def _embed_avatars(
    ws: Any,
    df: pd.DataFrame,
    avatar_urls: list[str | None],
    session: requests.Session,
    max_side_px: int,
) -> None:
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter

    if "頭像" not in df.columns or len(avatar_urls) != len(df):
        return

    col_idx = list(df.columns).index("頭像") + 1
    letter = get_column_letter(col_idx)

    for i, url in enumerate(avatar_urls):
        if not url:
            continue
        raw = _download_avatar_image(session, url)
        if not raw:
            continue
        try:
            pil = PILImage.open(io.BytesIO(raw))
            pil = pil.convert("RGBA")
            pil.thumbnail((max_side_px, max_side_px), PILImage.Resampling.LANCZOS)
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            buf.seek(0)
            tmp_path: str | None = None
            try:
                xlimg = XLImage(buf)
            except Exception:
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    pil.save(tmp, format="PNG")
                    tmp_path = tmp.name
                xlimg = XLImage(tmp_path)

            row_excel = i + 2
            xlimg.width = pil.width
            xlimg.height = pil.height
            ws.add_image(xlimg, f"{letter}{row_excel}")
            ws.row_dimensions[row_excel].height = max(
                ws.row_dimensions[row_excel].height or 15,
                min(pil.height * 0.78, 120),
            )
            if tmp_path:
                try:
                    os.unlink(tmp_path)
                except OSError:
                    pass
        except OSError:
            continue


def _apply_font_microsoft_jhenghei_and_header_fill(ws: Any, df: pd.DataFrame) -> None:
    """全表微軟正黑體；第一列標題加底色、粗體。"""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Font, PatternFill

    body_font = Font(name="Microsoft JhengHei", size=11)
    header_font = Font(name="Microsoft JhengHei", size=11, bold=True)
    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid",
    )

    max_r = max(1, ws.max_row or 1)
    max_c = max(1, ws.max_column or len(df.columns))

    for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=max_c):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.row == 1:
                cell.font = header_font
                cell.fill = header_fill
            else:
                cell.font = body_font


def _write_sessions_sheet_openpyxl(
    wb: Any,
    df: pd.DataFrame,
    cfg: dict[str, Any],
    avatar_urls: list[str | None],
    session: requests.Session,
) -> None:
    from openpyxl import Workbook

    max_px = int(cfg.get("avatar_max_px") or 52)
    if "sessions" in wb.sheetnames:
        del wb["sessions"]
    ws = wb.create_sheet("sessions", 0)
    if not df.empty:
        for j, col in enumerate(df.columns, start=1):
            ws.cell(1, j, value=col)
        for i, row in enumerate(df.itertuples(index=False), start=2):
            for j, val in enumerate(row, start=1):
                ws.cell(i, j, value=val)
    _autofit_worksheet_columns(ws, df)
    if cfg.get("merge_session_cells", True):
        _merge_session_info_cells(ws, df)
    _embed_avatars(ws, df, avatar_urls, session, max_px)
    _apply_font_microsoft_jhenghei_and_header_fill(ws, df)
    if isinstance(wb, Workbook) and "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        try:
            wb.remove(wb["Sheet"])
        except (ValueError, KeyError):
            pass


def _append_history_rows_openpyxl(ws: Any, append_df: pd.DataFrame) -> None:
    """在標題列下方插入新列（舊列下移），不調整欄寬。"""
    n = len(append_df)
    if n <= 0:
        return
    ws.insert_rows(2, amount=n)
    cols = list(append_df.columns)
    for i in range(n):
        for j, c in enumerate(cols, start=1):
            val = append_df.iloc[i][c]
            if pd.isna(val):
                val = ""
            ws.cell(2 + i, j, value=val)


def write_excel(
    df: pd.DataFrame,
    cfg: dict[str, Any],
    avatar_urls: list[str | None],
    session: requests.Session,
    *,
    history_append_df: pd.DataFrame | None = None,
    history_reset: bool = False,
) -> Path:
    from openpyxl import Workbook, load_workbook

    out = Path(__file__).resolve().parent / cfg["output_excel"]
    history_sheet_cols = list(HISTORY_SHEET_COLUMNS)
    touch_history = history_reset or (
        history_append_df is not None and not history_append_df.empty
    )

    if out.is_file() and not touch_history:
        wb = load_workbook(out)
        _write_sessions_sheet_openpyxl(wb, df, cfg, avatar_urls, session)
        wb.save(out)
        return out

    if out.is_file():
        wb = load_workbook(out)
    else:
        wb = Workbook()

    _write_sessions_sheet_openpyxl(wb, df, cfg, avatar_urls, session)

    if history_reset:
        if HISTORY_SHEET_NAME in wb.sheetnames:
            del wb[HISTORY_SHEET_NAME]
        hs = wb.create_sheet(HISTORY_SHEET_NAME)
        for j, h in enumerate(history_sheet_cols, start=1):
            hs.cell(1, j, value=h)
    elif history_append_df is not None and not history_append_df.empty:
        if HISTORY_SHEET_NAME in wb.sheetnames:
            hs = wb[HISTORY_SHEET_NAME]
        else:
            hs = wb.create_sheet(HISTORY_SHEET_NAME)
            for j, h in enumerate(history_sheet_cols, start=1):
                hs.cell(1, j, value=h)
        hist_df = history_append_df.reindex(columns=history_sheet_cols)
        _append_history_rows_openpyxl(hs, hist_df)

    wb.save(out)
    return out


def _load_baseline(cfg: dict[str, Any], sh: Any) -> dict[str, dict[str, Any]]:
    from google_sheets_export import google_sheets_enabled, load_baseline_from_sheet

    if google_sheets_enabled(cfg) and sh is not None:
        return load_baseline_from_sheet(sh)
    return load_last_scan_state()


def _save_baseline(
    cfg: dict[str, Any], sh: Any, state: dict[str, dict[str, Any]]
) -> None:
    from google_sheets_export import google_sheets_enabled, save_baseline_to_sheet

    if google_sheets_enabled(cfg) and sh is not None:
        save_baseline_to_sheet(sh, state)
        if LAST_SCAN_STATE_PATH.is_file():
            try:
                LAST_SCAN_STATE_PATH.unlink()
            except OSError:
                pass
    else:
        save_last_scan_state(state)


def _clear_history(cfg: dict[str, Any], sh: Any, out_path: Path) -> bool:
    """清空 掃描歷史 分頁／工作表，僅保留標題列。回傳是否清空成功。"""
    from google_sheets_export import (
        clear_history_sheet,
        google_sheets_enabled,
    )

    if google_sheets_enabled(cfg) and sh is not None:
        clear_history_sheet(sh)
        return True
    return False


def main() -> None:
    from google_sheets_export import google_sheets_enabled, open_spreadsheet, write_google_sheets

    parser = argparse.ArgumentParser(description="FitBook 場次爬蟲")
    parser.add_argument(
        "--reset-baseline",
        action="store_true",
        help="清空比對基準與掃描歷史，本次掃描結果作為新的基準（不寫入新增/刪除列）",
    )
    parser.add_argument(
        "--full",
        action="store_true",
        help="略過快掃，強制完整掃描（抓會員頁）",
    )
    parser.add_argument(
        "--quick-only",
        action="store_true",
        help="僅執行快掃並結束（測試用）",
    )
    args = parser.parse_args()

    cfg = load_config()

    if args.quick_only:
        need_deep, reason = run_quick_check(cfg)
        print(
            f"快掃：{reason}（{'需深度掃描' if need_deep else '略過深度掃描'}）"
        )
        return

    if (
        not args.reset_baseline
        and not args.full
        and quick_scan_enabled(cfg)
    ):
        need_deep, reason = run_quick_check(cfg)
        if not need_deep:
            print(f"快掃：{reason}，略過深度掃描")
            return
        print(f"快掃：{reason}，啟動深度掃描")

    try:
        df, avatar_urls, session = run_once()
    except (ConnectionError, Timeout, ChunkedEncodingError, requests.HTTPError) as e:
        print(f"連線失敗，本次掃描略過：{e}")
        return

    out = Path(__file__).resolve().parent / cfg["output_excel"]

    history_cols = list(HISTORY_COLUMNS)
    current_state = session_state_from_dataframe(df, avatar_urls)
    tz = _scan_timezone(cfg)
    today = datetime.now(tz).date()
    current_state = filter_state_active_sessions(current_state, today)
    placeholder_sessions = _count_placeholder_sessions(df)
    sheets_on = google_sheets_enabled(cfg)
    sh = open_spreadsheet(cfg) if sheets_on else None

    history_append_df: pd.DataFrame | None = None
    prev_baseline_avatar_map: dict[tuple[str, str], str] = {}
    history_reset = False
    info = ""
    added_n = removed_n = 0

    if args.reset_baseline:
        if placeholder_sessions > 0:
            info = (
                f"reset 終止：本次有 {placeholder_sessions} 個場次無法解析會員，"
                f"請先確認 Cookie 與網路後重試"
            )
        else:
            history_reset = True
            _save_baseline(cfg, sh, current_state)
            info = f"已重設比對基準，「{HISTORY_SHEET_NAME}」已清空"
    else:
        baseline = _load_baseline(cfg, sh)
        baseline_compare = filter_state_active_sessions(baseline, today)
        has_baseline = bool(baseline_compare)

        if placeholder_sessions > 0:
            info = (
                f"本次有 {placeholder_sessions} 個場次無法解析會員，"
                f"略過掃描歷史與基準更新"
            )
        elif not has_baseline:
            _save_baseline(cfg, sh, current_state)
            info = "首次建立比對基準（無新增/刪除）"
        elif snapshot_has_member_change(baseline_compare, current_state):
            chunk = build_history_diff_chunk(
                df,
                baseline_compare,
                current_state,
                history_cols,
                today=today,
                avatar_urls=avatar_urls,
            )
            if not chunk.empty and "異動類型" in chunk.columns:
                added_n = int((chunk["異動類型"] == "新增").sum())
                removed_n = int((chunk["異動類型"] == "刪除").sum())
                total_curr = sum(
                    len((v or {}).get("members") or [])
                    for v in current_state.values()
                )
                total_prev = sum(
                    len((v or {}).get("members") or [])
                    for v in baseline_compare.values()
                )
                if (
                    removed_n >= 8
                    and removed_n > added_n
                    and (total_prev == 0 or total_curr / max(total_prev, 1) < 0.5)
                ):
                    info = (
                        f"異動異常（刪除 {removed_n}、新增 {added_n}，"
                        f"目前/基準會員數 {total_curr}/{total_prev}），"
                        f"略過寫入掃描歷史且不更新基準"
                    )
                else:
                    history_append_df = _sort_dataframe_by_session_date(chunk)
                    from google_sheets_export import baseline_member_avatar_map

                    prev_baseline_avatar_map = baseline_member_avatar_map(
                        baseline_compare
                    )
                    _save_baseline(cfg, sh, current_state)
                    info = (
                        f"人員異動 {len(chunk)} 列"
                        f"（新增 {added_n}、刪除 {removed_n}）"
                    )
            else:
                _save_baseline(cfg, sh, current_state)
                info = "無有效異動，未追加歷史"
        else:
            info = "與基準無異動，未追加歷史"

    if sheets_on:
        url = write_google_sheets(
            df,
            avatar_urls,
            cfg,
            history_append_df=history_append_df,
            history_reset=history_reset,
            spreadsheet=sh,
        )
        path_msg = url
    else:
        path_msg = str(
            write_excel(
                df,
                cfg,
                avatar_urls,
                session,
                history_append_df=history_append_df,
                history_reset=history_reset,
            )
        )

    notify_parts: list[str] = []
    if history_append_df is not None and not history_append_df.empty:
        sheet_url = path_msg if sheets_on else None
        notify_diff_df = history_append_df
        try:
            from google_sheets_export import enrich_diff_df_avatars, load_sessions_avatar_map

            scrape_avatar_map = _avatar_url_map_from_df(df, avatar_urls)
            sessions_avatar_map = (
                load_sessions_avatar_map(sh)
                if sheets_on and sh is not None
                else scrape_avatar_map
            )
            notify_diff_df = enrich_diff_df_avatars(
                history_append_df,
                sessions_map=sessions_avatar_map,
                baseline_avatar_map=prev_baseline_avatar_map,
                scrape_map=scrape_avatar_map,
            )
        except Exception:
            notify_diff_df = history_append_df
        try:
            from telegram_notify import maybe_send_telegram_diff

            tg_result = maybe_send_telegram_diff(
                cfg,
                notify_diff_df,
                sheet_url=sheet_url,
                http_session=session,
            )
            if tg_result:
                notify_parts.append(tg_result)
        except Exception as e:
            notify_parts.append(f"Telegram 通知失敗：{e}")
        try:
            from line_notify import maybe_send_line_diff

            line_result = maybe_send_line_diff(
                cfg, history_append_df, sheet_url=sheet_url
            )
            if line_result:
                notify_parts.append(line_result)
        except Exception as e:
            notify_parts.append(f"LINE 通知失敗：{e}")

    notify_info = f"；{'；'.join(notify_parts)}" if notify_parts else ""
    print(f"寫入: {path_msg}，目前筆數: {len(df)}；{info}{notify_info}")


if __name__ == "__main__":
    main()
